from __future__ import annotations

import csv
import io
from pathlib import Path

from docx import Document
from fastapi import UploadFile
from loguru import logger
from openpyxl import load_workbook
from PyPDF2 import PdfReader

from services.document_cleaner import remove_unwanted_sections

# MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MiB per upload
MAX_FILE_SIZE = 500 * 1024 * 1024  # 500 MiB per upload
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx"}


class DocumentExtractionError(ValueError):
    """Raised when an uploaded file cannot safely be converted to text."""


async def extract_document_text(upload: UploadFile) -> str:
    filename = upload.filename or ""
    extension = Path(filename).suffix.lower()
    logger.info("Extracting document text for file: '{}' (extension: '{}')", filename, extension)

    if extension not in SUPPORTED_EXTENSIONS:
        accepted = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        logger.error("Unsupported file type for '{}'. Accepted types: {}", filename, accepted)
        raise DocumentExtractionError(f"Unsupported file type for '{filename}'. Accepted: {accepted}")

    data = await upload.read(MAX_FILE_SIZE + 1)
    # if not data:
    #     logger.error("File '{}' is empty", filename)
    #     raise DocumentExtractionError(f"'{filename}' is empty")
    # if len(data) > MAX_FILE_SIZE:
    #     logger.error("File '{}' exceeds 10 MiB limit (size: {} bytes)", filename, len(data))
    #     raise DocumentExtractionError(f"'{filename}' exceeds the 10 MiB upload limit")

    if len(data) > MAX_FILE_SIZE:
        logger.error(
            "File '{}' exceeds 500 MiB limit (size: {} bytes)",
            filename,
            len(data),
        )
        raise DocumentExtractionError(
            f"'{filename}' exceeds the 500 MiB upload limit"
        )

    try:
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                logger.error("File '{}' is password protected", filename)
                raise DocumentExtractionError(f"'{filename}' is password protected")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            document = Document(io.BytesIO(data))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
        elif extension == ".xlsx":
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for worksheet in workbook.worksheets:
                lines.append(f"[Sheet: {worksheet.title}]")
                for row in worksheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        lines.append(" | ".join(values))
            text = "\n".join(lines)
            workbook.close()
        elif extension == ".csv":
            decoded = data.decode("utf-8-sig")
            text = "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(decoded)))
        else:
            text = data.decode("utf-8")
    except DocumentExtractionError:
        raise
    except Exception as exc:
        logger.exception("Could not read file '{}'", filename)
        raise DocumentExtractionError(f"Could not read '{filename}'") from exc

    text = text.strip()
    if not text:
        logger.error("No readable text found in '{}'", filename)
        raise DocumentExtractionError(f"No readable text was found in '{filename}'")

    logger.info("Raw text extracted from '{}' (length: {} chars). Applying section removal...", filename, len(text))

    text = remove_unwanted_sections(text)
    if not text:
        logger.error("No readable text remained in '{}' after section removal", filename)
        raise DocumentExtractionError(f"No readable text remained in '{filename}' after section removal")

    logger.info("Successfully extracted and cleaned document text for '{}'", filename)
    return text
